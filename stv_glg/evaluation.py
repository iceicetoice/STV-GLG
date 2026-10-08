from __future__ import annotations

from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import torch

from .data import make_batch, preceding_batch


def metrics(actual, prediction) -> dict:
    actual, prediction = np.asarray(actual, dtype=float), np.asarray(prediction, dtype=float)
    if actual.shape != prediction.shape or actual.size == 0:
        raise ValueError("Metrics require nonempty arrays of identical shape.")
    if not np.isfinite(actual).all() or not np.isfinite(prediction).all():
        raise ValueError("Metrics require finite values.")
    error = prediction - actual
    mse = float(np.mean(error ** 2))
    sst = float(((actual - actual.mean()) ** 2).sum())
    denominator = float(np.abs(actual).sum())
    return {"MAE_km2": float(np.abs(error).mean()), "MSE_km4": mse,
            "RMSE_km2": float(np.sqrt(mse)), "R2": 1 - float((error ** 2).sum()) / sst if sst > 1e-12 else np.nan,
            "WAPE_percent": 100 * float(np.abs(error).sum()) / denominator if denominator > 1e-12 else np.nan,
            "n_dates": int(actual.size)}


@torch.inference_mode()
def evaluate(model, seasons, years, normalizer, config, device, include_updates=False):
    was_training = model.training
    model.eval()
    records, updates = [], []
    for year in years:
        season = seasons[year]
        for n in config["eval_prefixes"]:
            if n >= config["max_days"] or season.target_mask[n:].sum() < 3:
                continue
            batch = make_batch([season], n, normalizer, device)
            result = model(batch)
            predictions = {"STV-GLGFormer": result["prediction"][0].cpu().numpy() * normalizer.area_scale}
            future_mask = season.target_mask[n:]
            for method, prediction in predictions.items():
                for index in np.flatnonzero(future_mask):
                    records.append({
                        "year": year, "prefix_days": n, "issue_date": season.dates[n - 1],
                        "date": season.dates[n + index], "lead_days": int(index + 1),
                        "method": method, "observed_area_km2": float(season.target[n + index]),
                        "predicted_area_km2": float(prediction[index]),
                        "original_observation": bool(season.observed_mask[n + index]),
                    })
            if include_updates and n > config["min_prefix"]:
                previous = model(preceding_batch(batch, config["min_prefix"]))["prediction"][0, 1:].cpu().numpy() * normalizer.area_scale
                current = predictions["STV-GLGFormer"]
                actual = season.target[n:][future_mask]
                cur_error = float(np.mean((current[future_mask] - actual) ** 2))
                pre_error = float(np.mean((previous[future_mask] - actual) ** 2))
                updates.append({"year": year, "prefix_days": n, "current_mse_km4": cur_error,
                                "previous_mse_km4": pre_error, "increase_km4": cur_error - pre_error,
                                "excess_km4": max(cur_error - pre_error - config["training"]["gamma_km4"], 0),
                                "mean_absolute_revision_km2": float(np.abs(current[future_mask] - previous[future_mask]).mean())})
    model.train(was_training)
    frame = pd.DataFrame.from_records(records)
    if frame.empty:
        raise ValueError("No evaluable prefixes/years.")
    per_prefix = []
    for (year, n, method), group in frame.groupby(["year", "prefix_days", "method"]):
        per_prefix.append({"year": year, "prefix_days": n, "method": method,
                           **metrics(group.observed_area_km2, group.predicted_area_km2)})
    return frame, pd.DataFrame(per_prefix), pd.DataFrame(updates)


def summarize(forecasts: pd.DataFrame, per_prefix: pd.DataFrame) -> pd.DataFrame:
    rows = []
    fields = ["MAE_km2", "MSE_km4", "RMSE_km2", "R2", "WAPE_percent"]
    for method, group in forecasts.groupby("method"):
        rows.append({"aggregation": "pooled_year_prefix_date", "method": method,
                     **metrics(group.observed_area_km2, group.predicted_area_km2)})
        part = per_prefix.loc[per_prefix.method.eq(method)]
        # Years have equal weight, and each checkpoint has equal weight within year.
        year_means = part.groupby("year")[fields].mean()
        rows.append({"aggregation": "macro_year_then_prefix", "method": method,
                     **year_means.mean().to_dict(), "n_dates": int(len(group))})
        original = group.loc[group.original_observation]
        if not original.empty:
            rows.append({"aggregation": "pooled_original_observations_only", "method": method,
                         **metrics(original.observed_area_km2, original.predicted_area_km2)})
    return pd.DataFrame(rows)


def disappearance_date(dates, prediction, historical_peak, threshold=5.0, consecutive=3, require_post_peak=True):
    prediction = np.asarray(prediction)
    if require_post_peak and historical_peak < threshold and prediction.max() < threshold:
        return None
    # Do not call a low-area pre-bloom interval disappearance.
    start = int(prediction.argmax()) if require_post_peak and prediction.max() >= threshold else 0
    run = 0
    for i in range(start, len(prediction)):
        run = run + 1 if prediction[i] < threshold else 0
        if run >= consecutive:
            return pd.Timestamp(dates[i])
    return None


def plot_prediction(path, season, n, prediction, stop):
    from .data import historical_area
    from matplotlib.dates import DateFormatter
    fig, ax = plt.subplots(figsize=(11, 4.8))
    history = historical_area(season, n)
    ax.plot(season.dates[:n], history, color="#222222", lw=1.6, label="Historical area")
    future_dates = season.dates[n:]
    valid = season.target_mask[n:]
    if valid.any():
        ax.plot(future_dates[valid], season.target[n:][valid], color="#777777", lw=1.2,
                label="Available future reference (evaluation only)")
    ax.plot(future_dates, prediction, color="#0072B2", lw=1.8, label="STV-GLGFormer")
    ax.axvline(season.dates[n - 1], color="#777777", linestyle="--", lw=0.9, label="Forecast issue")
    if stop is not None:
        ax.axvline(stop, color="#D55E00", linestyle=":", lw=1.2, label="Operational disappearance")
    ax.set(title=f"{season.year} | issue: {season.dates[n - 1]:%Y-%m-%d} | prefix: {n} days",
           ylabel="Coverage area (km$^2$)", xlabel="Date")
    ax.xaxis.set_major_formatter(DateFormatter("%m-%d"))
    ax.grid(axis="y", alpha=0.2)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, fontsize=9)
    fig.tight_layout()
    fig.savefig(path, dpi=180)
    plt.close(fig)


def export_evaluation(output_dir, forecasts, per_prefix, updates, seasons, config):
    output = Path(output_dir)
    output.mkdir(parents=True, exist_ok=True)
    summary = summarize(forecasts, per_prefix)
    forecasts.to_csv(output / "forecasts_long.csv", index=False, encoding="utf-8-sig")
    per_prefix.to_csv(output / "metrics_by_year_prefix.csv", index=False, encoding="utf-8-sig")
    summary.to_csv(output / "metrics_summary.csv", index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(output / "test_forecasts_and_metrics.xlsx", engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="summary", index=False)
        per_prefix.to_excel(writer, sheet_name="year_prefix_metrics", index=False)
        forecasts.to_excel(writer, sheet_name="all_future_predictions", index=False)
        if not updates.empty:
            updates.to_excel(writer, sheet_name="adjacent_prefix_updates", index=False)
        model_rows = forecasts.loc[forecasts.method.eq("STV-GLGFormer")]
        for (year, n), group in model_rows.groupby(["year", "prefix_days"]):
            original = seasons[year].original.copy()
            keyed = group.set_index("date").predicted_area_km2
            original["STV-GLGFormer预测面积（km²）"] = pd.to_datetime(original[config["date_column"]]).map(keyed)
            original.to_excel(writer, sheet_name=f"{year}_N{n}", index=False)
    figure_dir = output / "figures"
    figure_dir.mkdir(exist_ok=True)
    for year in sorted(forecasts.year.unique()):
        groups = forecasts.loc[forecasts.year.eq(year)]
        prefixes = sorted(groups.prefix_days.unique())
        fig, axes = plt.subplots(len(prefixes), 1, figsize=(11, 3.2 * len(prefixes)), squeeze=False)
        for ax, n in zip(axes[:, 0], prefixes):
            data = groups.loc[groups.prefix_days.eq(n)]
            truth = data.loc[data.method.eq("STV-GLGFormer")]
            ax.plot(truth.date, truth.observed_area_km2, color="#222222", lw=1.6, label="Area reference")
            for method, sub in data.groupby("method"):
                ax.plot(sub.date, sub.predicted_area_km2, lw=1.5 if method == "STV-GLGFormer" else 0.9,
                        alpha=1 if method == "STV-GLGFormer" else 0.65, label=method)
            ax.set_title(f"{year} | observed prefix: {n} days | issue: {truth.issue_date.iloc[0]:%m-%d}", fontsize=11)
            ax.set_ylabel("Coverage area (km$^2$)")
            ax.grid(axis="y", alpha=0.2)
            from matplotlib.dates import DateFormatter
            ax.xaxis.set_major_formatter(DateFormatter("%m-%d"))
            ax.spines[["top", "right"]].set_visible(False)
        axes[0, 0].legend(fontsize=8, ncol=3, frameon=False)
        fig.tight_layout()
        fig.savefig(figure_dir / f"{year}_prefix_forecasts.png", dpi=170)
        plt.close(fig)
    for book in [output / "test_forecasts_and_metrics.xlsx"]:
        from openpyxl import load_workbook
        workbook = load_workbook(book)
        for sheet in workbook:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                sheet.column_dimensions[cell.column_letter].width = 24
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, (pd.Timestamp, __import__("datetime").datetime)):
                        cell.number_format = "yyyy-mm-dd"
        workbook.save(book)
    return summary
